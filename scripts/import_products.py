#!/usr/bin/env python3
"""
Import or update Shopify products from the dub SEO product metadata spreadsheet.

Upsert uses Admin GraphQL productSet with ProductSetIdentifiers { handle }.

Environment:
  SHOPIFY_SHOP          Store domain, e.g. your-store.myshopify.com (required unless --shop)
  SHOPIFY_ACCESS_TOKEN  Custom app Admin API access token with read_products, write_products (required)

Optional:
  SHOPIFY_API_VERSION   Default: 2025-10
  PLACEHOLDER_PRICE     Default: 0.01 (string or number; used for every variant)

Install dependencies (use python3 -m pip if `pip` is not on your PATH):
  python3 -m pip install -r requirements-scripts.txt

Example:
  export SHOPIFY_SHOP=your-store.myshopify.com
  export SHOPIFY_ACCESS_TOKEN=shpat_...
  python3 scripts/import_products.py --xlsx assets/dub_SEO_Product_Metadata.xlsx
  python3 scripts/import_products.py --xlsx assets/dub_SEO_Product_Metadata.xlsx --dry-run
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from typing import Any

import requests
from requests import HTTPError

try:
    from openpyxl import load_workbook
except ImportError:
    print(
        "Missing dependency: python3 -m pip install -r requirements-scripts.txt",
        file=sys.stderr,
    )
    raise

API_VERSION_DEFAULT = os.environ.get("SHOPIFY_API_VERSION", "2025-10")

PRODUCT_SET_MUTATION = """
mutation UpsertProduct($input: ProductSetInput!, $identifier: ProductSetIdentifiers) {
  productSet(input: $input, identifier: $identifier) {
    product {
      id
      handle
      title
      variants(first: 100) {
        nodes {
          id
          title
          price
        }
      }
    }
    userErrors {
      field
      message
      code
    }
  }
}
"""

PRODUCT_LOOKUP_QUERY = """
query ProductByHandle($q: String!) {
  products(first: 1, query: $q) {
    nodes {
      id
      handle
    }
  }
}
"""


def normalize_shop(shop: str) -> str:
    shop = shop.strip().lower()
    if not shop.endswith(".myshopify.com"):
        if "." in shop:
            return shop
        return f"{shop}.myshopify.com"
    return shop


def normalize_title(name: str) -> str:
    if not name:
        return ""
    s = " ".join(name.split())
    return s.strip()


def normalize_handle(slug: str | None, sku_handle: str | None) -> str:
    raw = (slug or "").strip() or (sku_handle or "").strip()
    if not raw:
        return ""
    h = raw.lower()
    h = re.sub(r"[^a-z0-9-]+", "-", h)
    h = re.sub(r"-+", "-", h).strip("-")
    return h


def parse_colors(cell: str | None) -> list[str]:
    if not cell or not str(cell).strip():
        return ["Default"]
    parts = [p.strip() for p in str(cell).split(",")]
    out = [p for p in parts if p]
    return out or ["Default"]


def slug_for_sku_fragment(text: str) -> str:
    t = text.lower().strip()
    t = re.sub(r"[^a-z0-9]+", "-", t)
    return re.sub(r"-+", "-", t).strip("-") or "variant"


def build_option_values(
    option_name: str, colors: list[str]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    values = [{"name": c} for c in colors]
    variants = []
    for c in colors:
        variants.append(
            {
                "optionValues": [{"optionName": option_name, "name": c}],
                "price": None,  # set by caller
                "sku": None,  # set by caller
            }
        )
    return values, variants


def _http_error_message(shop: str, api_version: str, err: HTTPError) -> str:
    resp = err.response
    status = resp.status_code if resp is not None else 0
    snippet = ""
    if resp is not None and resp.text:
        snippet = resp.text.strip()[:400]
    lines = [
        f"HTTP {status} from Shopify Admin API ({shop}, API {api_version}).",
    ]
    if status == 401:
        lines.extend(
            [
                "Unauthorized: the Admin API token was rejected.",
                "Fix: In Shopify Admin → Settings → Apps → Develop apps → your app,",
                "install the app on this store and copy a fresh Admin API access token",
                "(starts with shpat_). Ensure scopes include read_products and write_products.",
                "The token must be for this exact shop; theme tokens and Storefront API tokens will not work.",
            ]
        )
    elif status == 403:
        lines.append(
            "Forbidden: token may be valid but missing required API scopes for this operation."
        )
    elif status == 404:
        lines.append(
            "Not found: check SHOPIFY_SHOP (e.g. your-store.myshopify.com) and API version."
        )
    else:
        lines.append("See response body below.")
    if snippet:
        lines.append(f"Response: {snippet}")
    return "\n".join(lines)


def graphql_request(
    shop: str,
    token: str,
    api_version: str,
    query: str,
    variables: dict[str, Any] | None = None,
) -> dict[str, Any]:
    token = (token or "").strip()
    if not token:
        raise RuntimeError(
            "Missing SHOPIFY_ACCESS_TOKEN (or --token). Admin API access token is required."
        )
    url = f"https://{shop}/admin/api/{api_version}/graphql.json"
    headers = {
        "Content-Type": "application/json",
        "X-Shopify-Access-Token": token,
    }
    body: dict[str, Any] = {"query": query}
    if variables is not None:
        body["variables"] = variables
    r = requests.post(url, headers=headers, json=body, timeout=120)
    try:
        r.raise_for_status()
    except HTTPError as e:
        raise RuntimeError(_http_error_message(shop, api_version, e)) from e
    data = r.json()
    if "errors" in data and data["errors"]:
        raise RuntimeError(f"GraphQL errors: {json.dumps(data['errors'], indent=2)}")
    return data.get("data") or {}


def product_exists(shop: str, token: str, api_version: str, handle: str) -> bool:
    # Shopify search syntax: handle:exact-handle
    q = f"handle:{handle}"
    data = graphql_request(
        shop,
        token,
        api_version,
        PRODUCT_LOOKUP_QUERY,
        {"q": q},
    )
    nodes = (data.get("products") or {}).get("nodes") or []
    return len(nodes) > 0


def read_rows(xlsx_path: str) -> tuple[list[str], list[list[Any]]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(c).strip() if c is not None else "" for c in next(rows_iter)]
    except StopIteration:
        return [], []
    data: list[list[Any]] = []
    for row in rows_iter:
        data.append(list(row))
    return header, data


def row_dict(header: list[str], row: list[Any]) -> dict[str, Any]:
    d: dict[str, Any] = {}
    for i, key in enumerate(header):
        if not key:
            continue
        d[key.strip()] = row[i] if i < len(row) else None
    return d


def build_product_set_input(
    row: dict[str, Any],
    option_name: str,
    placeholder_price: str,
    handle: str,
) -> dict[str, Any]:
    title = normalize_title(str(row.get("Product Name") or ""))
    desc = (row.get("Product Description (Brand Voice)") or "") or ""
    if isinstance(desc, str):
        desc_html = desc.strip()
    else:
        desc_html = str(desc)

    page_title = (row.get("Page Title (≤60 chars)") or row.get("Page Title") or "") or ""
    meta_desc = (
        row.get("Meta Description (≤155 chars)") or row.get("Meta Description") or ""
    ) or ""
    if not isinstance(page_title, str):
        page_title = str(page_title)
    if not isinstance(meta_desc, str):
        meta_desc = str(meta_desc)

    category = (row.get("Category") or "") or ""
    if not isinstance(category, str):
        category = str(category)
    product_type = category.strip() or None

    colors = parse_colors(row.get("Colors Available"))
    values, variant_templates = build_option_values(option_name, colors)

    metafields: list[dict[str, Any]] = []
    pk = row.get("Primary Keyword")
    if pk is not None and str(pk).strip():
        metafields.append(
            {
                "namespace": "custom",
                "key": "primary_keyword",
                "value": str(pk).strip(),
                "type": "single_line_text_field",
            }
        )
    sk = row.get("Secondary Keywords")
    if sk is not None and str(sk).strip():
        metafields.append(
            {
                "namespace": "custom",
                "key": "secondary_keywords",
                "value": str(sk).strip(),
                "type": "multi_line_text_field",
            }
        )
    alt = row.get("Image Alt Text")
    if alt is not None and str(alt).strip():
        metafields.append(
            {
                "namespace": "custom",
                "key": "image_alt_suggestion",
                "value": str(alt).strip(),
                "type": "multi_line_text_field",
            }
        )
    src = row.get("Source URL")
    if src is not None and str(src).strip():
        metafields.append(
            {
                "namespace": "custom",
                "key": "source_url",
                "value": str(src).strip(),
                "type": "single_line_text_field",
            }
        )

    variants: list[dict[str, Any]] = []
    for vt in variant_templates:
        color_name = vt["optionValues"][0]["name"]
        sku = f"{handle}-{slug_for_sku_fragment(color_name)}"
        price_val = float(placeholder_price)
        variants.append(
            {
                "optionValues": vt["optionValues"],
                "price": price_val,
                "sku": sku[:255],
            }
        )

    seo: dict[str, Any] = {}
    pt = page_title.strip()
    md = meta_desc.strip()
    if pt:
        seo["title"] = pt
    if md:
        seo["description"] = md

    inp: dict[str, Any] = {
        "title": title,
        "handle": handle,
        "descriptionHtml": desc_html,
        "status": "ACTIVE",
        "productOptions": [
            {
                "name": option_name,
                "position": 1,
                "values": values,
            }
        ],
        "variants": variants,
    }
    if seo:
        inp["seo"] = seo
    if product_type:
        inp["productType"] = product_type
    if metafields:
        inp["metafields"] = metafields
    return inp


def upsert_product(
    shop: str,
    token: str,
    api_version: str,
    handle: str,
    input_payload: dict[str, Any],
) -> dict[str, Any]:
    variables = {
        "identifier": {"handle": handle},
        "input": input_payload,
    }
    data = graphql_request(
        shop, token, api_version, PRODUCT_SET_MUTATION, variables
    )
    return data.get("productSet") or {}


def _errors_suggest_price_retry(errs: list[dict[str, Any]]) -> bool:
    for e in errs:
        msg = (e.get("message") or "").lower()
        if "price" in msg:
            return True
    return False


def run_upsert_with_price_fallback(
    shop: str,
    token: str,
    api_version: str,
    handle: str,
    rd: dict[str, Any],
    option_name: str,
    placeholder_price: str,
) -> dict[str, Any]:
    payload = build_product_set_input(rd, option_name, placeholder_price, handle)
    result = upsert_product(shop, token, api_version, handle, payload)
    errs = result.get("userErrors") or []
    if errs and _errors_suggest_price_retry(errs) and placeholder_price != "0.01":
        payload = build_product_set_input(rd, option_name, "0.01", handle)
        result = upsert_product(shop, token, api_version, handle, payload)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Shopify product import from Excel (upsert).")
    parser.add_argument(
        "--xlsx",
        default="assets/dub_SEO_Product_Metadata.xlsx",
        help="Path to workbook (default: assets/dub_SEO_Product_Metadata.xlsx)",
    )
    parser.add_argument("--shop", default=os.environ.get("SHOPIFY_SHOP", ""))
    parser.add_argument(
        "--token",
        default=os.environ.get("SHOPIFY_ACCESS_TOKEN", ""),
        help="Admin API token (default: SHOPIFY_ACCESS_TOKEN)",
    )
    parser.add_argument("--api-version", default=API_VERSION_DEFAULT)
    parser.add_argument(
        "--placeholder-price",
        default=os.environ.get("PLACEHOLDER_PRICE", "0.01"),
        help="Price for each variant (default 0.01)",
    )
    parser.add_argument(
        "--option-name",
        default="Color",
        help="Product option name for color variants (default: Color)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Parse spreadsheet and print actions without calling the API",
    )
    args = parser.parse_args()

    header, rows = read_rows(args.xlsx)
    if not header:
        print("No header row in spreadsheet.", file=sys.stderr)
        return 1

    shop = normalize_shop(args.shop) if args.shop else ""
    token = args.token.strip()

    if not args.dry_run:
        if not shop or not token:
            print(
                "Set SHOPIFY_SHOP and SHOPIFY_ACCESS_TOKEN (or use --shop / --token).",
                file=sys.stderr,
            )
            return 1

    seen_handles: set[str] = set()
    exit_code = 0

    for row in rows:
        rd = row_dict(header, row)
        name = normalize_title(str(rd.get("Product Name") or ""))
        if not name:
            continue

        handle = normalize_handle(
            rd.get("URL Slug"),
            rd.get("SKU / Handle"),
        )
        if not handle:
            print(f"SKIP (no handle): {name!r}", file=sys.stderr)
            continue
        if handle in seen_handles:
            print(f"ERROR duplicate handle in sheet: {handle}", file=sys.stderr)
            exit_code = 1
            continue
        seen_handles.add(handle)

        if args.dry_run:
            input_payload = build_product_set_input(
                rd, args.option_name, str(args.placeholder_price), handle
            )
            colors = parse_colors(rd.get("Colors Available"))
            print(
                f"[dry-run] handle={handle!r} title={input_payload['title']!r} "
                f"variants={len(colors)} option={args.option_name!r}"
            )
            continue

        existed = product_exists(shop, token, args.api_version, handle)
        action = "UPDATED" if existed else "CREATED"

        try:
            result = run_upsert_with_price_fallback(
                shop,
                token,
                args.api_version,
                handle,
                rd,
                args.option_name,
                str(args.placeholder_price),
            )
        except Exception as e:
            print(f"FAIL {handle}: {e}", file=sys.stderr)
            exit_code = 1
            continue

        errs = result.get("userErrors") or []
        if errs:
            print(
                f"FAIL {handle}: {json.dumps(errs, indent=2)}",
                file=sys.stderr,
            )
            exit_code = 1
            continue

        prod = result.get("product") or {}
        print(f"{action} handle={handle} id={prod.get('id')} title={prod.get('title')!r}")

    return exit_code


if __name__ == "__main__":
    sys.exit(main())
